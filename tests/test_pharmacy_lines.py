"""Pharmacy line-item parsing and export."""

from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from config import Settings
from models.document import (
    ExtractionFieldDefinition,
    ExtractionSchema,
    ExtractedFieldValue,
    FieldReviewStatus,
    ProcessingStatus,
    SourceDocument,
    StructuredExtractionResult,
)
from models.pharmacy import PharmacyLineItem
from providers.mapping import map_model_output_to_result
from providers.response_models import ModelFieldOutput, ModelLineItemOutput, ModelStructuringOutput
from services.excel_service import export_documents, preview_line_item_rows
from utils.pharmacy_lines import parse_product_lines_text


def test_parse_pipe_separated_product_lines() -> None:
    text = (
        "DRR | KETORAL DT TAB | STRP | 06/2029 | E2601499 | 60 | | 182.78 | "
        "139.26 | 8355.60 | 5.00 | 30049099"
    )
    items = parse_product_lines_text(text)
    assert len(items) == 1
    assert items[0].product_name == "KETORAL DT TAB"
    assert items[0].qty == "60"


def test_map_model_line_items() -> None:
    schema = ExtractionSchema(
        name="pharmacy_tax_invoice",
        fields=[
            ExtractionFieldDefinition(key="supplier_name", label="Supplier"),
            ExtractionFieldDefinition(key="product_lines", label="Lines"),
        ],
    )
    output = ModelStructuringOutput(
        fields=[
            ModelFieldOutput(key="supplier_name", value="VMA"),
            ModelFieldOutput(key="product_lines", value=""),
        ],
        line_items=[
            ModelLineItemOutput(product_name="VANTEJ PASTE", qty="18", amount="5307.48"),
        ],
    )
    result = map_model_output_to_result(schema, output)
    assert len(result.line_items) == 1
    assert result.line_items[0].product_name == "VANTEJ PASTE"


def test_map_falls_back_to_product_lines_text() -> None:
    schema = ExtractionSchema(
        name="pharmacy_tax_invoice",
        fields=[
            ExtractionFieldDefinition(key="product_lines", label="Lines"),
        ],
    )
    pipe_row = "DRR | CLOHEX ADS MW | 200ML | 06/2029 | X | 8 | | 100 | 90 | 720 | 5 | 123"
    output = ModelStructuringOutput(
        fields=[ModelFieldOutput(key="product_lines", value=pipe_row)],
    )
    result = map_model_output_to_result(schema, output)
    assert len(result.line_items) == 1
    assert result.line_items[0].product_name == "CLOHEX ADS MW"


def test_export_line_items_sheet(tmp_path: Path) -> None:
    schema = ExtractionSchema(
        name="pharmacy_tax_invoice",
        fields=[
            ExtractionFieldDefinition(key="supplier_name", label="Supplier"),
            ExtractionFieldDefinition(key="invoice_number", label="Inv"),
            ExtractionFieldDefinition(key="net_amount", label="Net"),
        ],
    )
    doc = SourceDocument(
        id=uuid4(),
        source_filename="inv.jpg",
        mime_type="image/jpeg",
        file_size_bytes=100,
        status=ProcessingStatus.READY_FOR_EXPORT,
        extraction=StructuredExtractionResult(
            fields=[
                ExtractedFieldValue(
                    key="supplier_name",
                    label="Supplier",
                    value="VMA",
                    review_status=FieldReviewStatus.CONFIRMED,
                ),
                ExtractedFieldValue(
                    key="invoice_number",
                    label="Inv",
                    value="ZR06394",
                    review_status=FieldReviewStatus.CONFIRMED,
                ),
                ExtractedFieldValue(
                    key="net_amount",
                    label="Net",
                    value="27078.00",
                    review_status=FieldReviewStatus.CONFIRMED,
                ),
            ],
            line_items=[
                PharmacyLineItem(product_name="A", qty="1", amount="10"),
                PharmacyLineItem(product_name="B", qty="2", amount="20"),
            ],
        ),
    )
    settings = Settings(_env_file=None)
    target = tmp_path / "out.xlsx"
    result = export_documents(target, [doc], schema, settings, mode="new")
    assert result.line_item_rows_written == 2
    wb = load_workbook(target)
    sheet = wb[settings.excel_line_items_sheet]
    assert sheet.cell(2, 9).value == "A"
    assert sheet.max_row == 3

    preview = preview_line_item_rows([doc])
    assert len(preview) == 2
