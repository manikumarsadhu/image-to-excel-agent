"""Excel preview and export for extracted documents."""

from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

from config import Settings
from models.document import (
    ExtractionSchema,
    ProcessingStatus,
    SourceDocument,
)
from models.pharmacy import PharmacyLineItem

META_COLUMNS: tuple[str, ...] = (
    "document_id",
    "source_filename",
    "processing_status",
    "export_review_status",
    "schema_name",
    "uploaded_at_utc",
    "ocr_engine",
    "ai_provider",
    "fingerprint",
    "error_message",
    "ocr_text",
)

EXPORTABLE_STATUSES: frozenset[ProcessingStatus] = frozenset(
    {
        ProcessingStatus.READY_FOR_EXPORT,
        ProcessingStatus.REVIEW_REQUIRED,
        ProcessingStatus.EXPORTED,
    }
)


class ExcelExportError(ValueError):
    """Export or append failed validation."""


LINE_ITEM_INVOICE_COLUMNS: tuple[str, ...] = (
    "document_id",
    "source_filename",
    "supplier_name",
    "invoice_number",
    "invoice_date",
    "billed_to",
    "net_amount",
)

LINE_ITEM_PRODUCT_COLUMNS: tuple[str, ...] = (
    "mfg",
    "product_name",
    "pack",
    "exp",
    "batch",
    "qty",
    "free",
    "mrp",
    "rate",
    "amount",
    "gst_pct",
    "hsn",
)


@dataclass(frozen=True)
class ExcelExportResult:
    path: Path
    rows_written: int
    sheet_name: str
    created_new_file: bool
    line_item_rows_written: int = 0
    line_items_sheet_name: str | None = None


def build_column_headers(schema: ExtractionSchema) -> list[str]:
    """Stable header order: metadata, then schema fields (+ review columns)."""
    headers: list[str] = list(META_COLUMNS)
    seen = set(headers)
    for field in schema.fields:
        value_col = field.key
        review_col = f"{field.key}_review"
        for col in (value_col, review_col):
            if col in seen:
                raise ExcelExportError(f"Duplicate Excel column name: {col}")
            seen.add(col)
            headers.append(col)
    return headers


def export_review_status(document: SourceDocument) -> str:
    if document.status == ProcessingStatus.REVIEW_REQUIRED or document.needs_review():
        return "needs_review"
    if document.status in (ProcessingStatus.READY_FOR_EXPORT, ProcessingStatus.EXPORTED):
        return "validated"
    return "other"


def document_to_row(document: SourceDocument, schema: ExtractionSchema) -> list[Any]:
    """Serialize one document to a row aligned with build_column_headers."""
    field_map = document.extraction.field_map() if document.extraction else {}
    uploaded = document.uploaded_at.astimezone(timezone.utc).replace(tzinfo=None)

    row: list[Any] = [
        str(document.id),
        document.source_filename,
        document.status.value,
        export_review_status(document),
        document.schema_name or schema.name,
        uploaded.isoformat(timespec="seconds"),
        document.ocr_engine or "",
        document.last_provider or "",
        document.fingerprint.value if document.fingerprint else "",
        document.error_message or "",
        document.ocr_text or "",
    ]

    for field_def in schema.fields:
        extracted = field_map.get(field_def.key)
        if extracted is None:
            row.extend(["", "missing"])
        else:
            row.extend([extracted.value or "", extracted.review_status.value])
    return row


def build_line_item_headers() -> list[str]:
    return list(LINE_ITEM_INVOICE_COLUMNS + LINE_ITEM_PRODUCT_COLUMNS)


def _extraction_field(document: SourceDocument, key: str) -> str:
    if document.extraction is None:
        return ""
    field = document.extraction.field_map().get(key)
    if field is None or field.value is None:
        return ""
    return str(field.value)


def line_item_to_row(document: SourceDocument, item: PharmacyLineItem) -> list[Any]:
    invoice_meta = [
        str(document.id),
        document.source_filename,
        _extraction_field(document, "supplier_name"),
        _extraction_field(document, "invoice_number"),
        _extraction_field(document, "invoice_date"),
        _extraction_field(document, "billed_to"),
        _extraction_field(document, "net_amount"),
    ]
    product = [
        item.mfg or "",
        item.product_name or "",
        item.pack or "",
        item.exp or "",
        item.batch or "",
        item.qty or "",
        item.free or "",
        item.mrp or "",
        item.rate or "",
        item.amount or "",
        item.gst_pct or "",
        item.hsn or "",
    ]
    return invoice_meta + product


def preview_line_item_rows(
    documents: Iterable[SourceDocument],
    *,
    statuses: frozenset[ProcessingStatus] | None = None,
) -> list[dict[str, Any]]:
    allowed = statuses or EXPORTABLE_STATUSES
    headers = build_line_item_headers()
    rows: list[dict[str, Any]] = []
    for document in documents:
        if document.status not in allowed or document.extraction is None:
            continue
        for item in document.extraction.line_items:
            if item.is_empty():
                continue
            rows.append(
                dict(zip(headers, line_item_to_row(document, item), strict=True))
            )
    return rows


def preview_rows(
    documents: Iterable[SourceDocument],
    schema: ExtractionSchema,
    *,
    statuses: frozenset[ProcessingStatus] | None = None,
) -> list[dict[str, Any]]:
    """Tabular preview for UI editing before export."""
    allowed = statuses or EXPORTABLE_STATUSES
    headers = build_column_headers(schema)
    rows: list[dict[str, Any]] = []
    for document in documents:
        if document.status not in allowed:
            continue
        if document.extraction is None:
            continue
        values = document_to_row(document, schema)
        rows.append(dict(zip(headers, values, strict=True)))
    return rows


def export_documents(
    target_path: Path | str,
    documents: list[SourceDocument],
    schema: ExtractionSchema,
    settings: Settings,
    *,
    mode: str = "new",
    overwrite_confirmed: bool = False,
    mark_exported: bool = True,
    statuses: frozenset[ProcessingStatus] | None = None,
) -> ExcelExportResult:
    """
    Write documents to Excel.

    mode:
      - new: create workbook (fails if file exists unless overwrite_confirmed)
      - append: add rows to existing sheet, preserving other sheets/rows
    """
    path = Path(target_path)
    if mode not in {"new", "append"}:
        raise ExcelExportError("mode must be 'new' or 'append'.")

    allowed = statuses or EXPORTABLE_STATUSES
    to_export = [
        doc
        for doc in documents
        if doc.status in allowed and doc.extraction is not None
    ]
    if not to_export:
        raise ExcelExportError("No exportable documents with extraction results.")

    headers = build_column_headers(schema)
    sheet_name = settings.excel_default_sheet

    if mode == "new":
        if path.exists() and not overwrite_confirmed:
            raise ExcelExportError(
                f"Refusing to overwrite existing workbook: {path.name}. "
                "Confirm overwrite or choose append mode."
            )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        _write_header_row(sheet, headers)
        next_row = 2
        created_new = True
    else:
        workbook, sheet, next_row, created_new = _open_for_append(path, sheet_name, headers)

    for document in to_export:
        sheet.append(document_to_row(document, schema))
        next_row += 1

    line_item_rows = 0
    line_items_sheet_name: str | None = None
    line_docs = [
        d
        for d in to_export
        if d.extraction and d.extraction.line_items
    ]
    if line_docs and schema.name == "pharmacy_tax_invoice":
        line_items_sheet_name = settings.excel_line_items_sheet
        li_headers = build_line_item_headers()
        if line_items_sheet_name in workbook.sheetnames:
            li_sheet = workbook[line_items_sheet_name]
            if li_sheet.max_row == 0 or _row_is_empty(li_sheet, 1):
                _write_header_row(li_sheet, li_headers)
            elif _read_header_row(li_sheet) != li_headers:
                raise ExcelExportError(
                    "Existing LineItems sheet headers do not match. Export to a new file."
                )
        else:
            li_sheet = workbook.create_sheet(line_items_sheet_name)
            _write_header_row(li_sheet, li_headers)
        for document in line_docs:
            assert document.extraction is not None
            for item in document.extraction.line_items:
                if item.is_empty():
                    continue
                li_sheet.append(line_item_to_row(document, item))
                line_item_rows += 1

    if mark_exported:
        for document in to_export:
            document.status = ProcessingStatus.EXPORTED
            document.error_message = None

    _atomic_save(workbook, path)
    return ExcelExportResult(
        path=path,
        rows_written=len(to_export),
        sheet_name=sheet_name,
        created_new_file=created_new,
        line_item_rows_written=line_item_rows,
        line_items_sheet_name=line_items_sheet_name,
    )


def _write_header_row(sheet: Worksheet, headers: list[str]) -> None:
    for idx, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=idx, value=header)


def _looks_like_xlsx(path: Path) -> bool:
    """True when the file has an Office Open XML (zip) signature."""
    try:
        with path.open("rb") as handle:
            header = handle.read(4)
    except OSError:
        return False
    return header[:2] == b"PK"


def _append_target_error(path: Path) -> ExcelExportError:
    suffix = path.suffix.lower()
    hint = (
        f"Cannot append to `{path}`: not a valid Excel .xlsx workbook. "
        "Use export mode **new** (check overwrite) to create a fresh file, "
        "or point append at an existing .xlsx you can open in Excel. "
    )
    if path.exists() and path.stat().st_size == 0:
        hint += "The file is empty (0 bytes)."
    elif suffix not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        hint += f"Expected a .xlsx path; got `{suffix or '(no extension)'}`. "
    elif path.exists() and not _looks_like_xlsx(path):
        hint += "The file is not a ZIP-based Excel workbook (wrong format or corrupted)."
    else:
        hint += "The file may be corrupted or still open in Excel — close it and retry."
    return ExcelExportError(hint)


def _load_workbook_or_raise(path: Path) -> WorkbookType:
    if path.exists() and path.stat().st_size == 0:
        raise _append_target_error(path)
    if path.exists() and not _looks_like_xlsx(path):
        raise _append_target_error(path)
    try:
        return load_workbook(path)
    except InvalidFileException as exc:
        raise _append_target_error(path) from exc


def _open_for_append(
    path: Path,
    sheet_name: str,
    expected_headers: list[str],
) -> tuple[WorkbookType, Worksheet, int, bool]:
    if path.exists():
        workbook = _load_workbook_or_raise(path)
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.create_sheet(sheet_name)
        if sheet.max_row == 0 or _row_is_empty(sheet, 1):
            _write_header_row(sheet, expected_headers)
            return workbook, sheet, 2, False
        existing = _read_header_row(sheet)
        if existing != expected_headers:
            raise ExcelExportError(
                "Existing workbook headers do not match the current schema. "
                "Export to a new file or align the schema."
            )
        return workbook, sheet, sheet.max_row + 1, False

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    _write_header_row(sheet, expected_headers)
    return workbook, sheet, 2, True


def _row_is_empty(sheet: Worksheet, row_index: int) -> bool:
    for cell in sheet[row_index]:
        if cell.value not in (None, ""):
            return False
    return True


def _read_header_row(sheet: Worksheet) -> list[str]:
    if sheet.max_row < 1:
        return []
    return [str(cell.value) if cell.value is not None else "" for cell in sheet[1]]


def _atomic_save(workbook: WorkbookType, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(path.name + ".tmp")
    try:
        workbook.save(tmp_path)
        os.replace(tmp_path, path)
    finally:
        if tmp_path.exists() and not path.exists():
            tmp_path.unlink(missing_ok=True)
