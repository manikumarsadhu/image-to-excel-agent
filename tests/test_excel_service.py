"""Excel export and preview tests."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from config import Settings
from models.document import (
    DocumentFingerprint,
    ExtractionFieldDefinition,
    ExtractionSchema,
    ExtractedFieldValue,
    FieldReviewStatus,
    ProcessingStatus,
    SourceDocument,
    StructuredExtractionResult,
)
from services.excel_service import (
    ExcelExportError,
    build_column_headers,
    export_documents,
    preview_rows,
)


@pytest.fixture
def settings() -> Settings:
    return Settings(_env_file=None, excel_default_sheet="ExtractedData")


@pytest.fixture
def schema() -> ExtractionSchema:
    return ExtractionSchema(
        name="receipt",
        fields=[
            ExtractionFieldDefinition(key="vendor", label="Vendor"),
            ExtractionFieldDefinition(key="total", label="Total"),
        ],
    )


def _doc(
    *,
    status: ProcessingStatus,
    vendor: str | None = "Café Müller",
    total: str | None = "₹500",
    review: FieldReviewStatus = FieldReviewStatus.CONFIRMED,
) -> SourceDocument:
    return SourceDocument(
        id=uuid4(),
        source_filename="scan.png",
        mime_type="image/png",
        file_size_bytes=1200,
        uploaded_at=datetime(2026, 1, 2, 12, 0, tzinfo=timezone.utc),
        ocr_text="CAFE MUELLER TOTAL 500",
        ocr_engine="paddleocr/en",
        schema_name="receipt",
        fingerprint=DocumentFingerprint(value="a" * 64),
        last_provider="gemini/test",
        status=status,
        extraction=StructuredExtractionResult(
            fields=[
                ExtractedFieldValue(
                    key="vendor",
                    label="Vendor",
                    value=vendor,
                    review_status=review if vendor else FieldReviewStatus.MISSING,
                ),
                ExtractedFieldValue(
                    key="total",
                    label="Total",
                    value=total,
                    review_status=FieldReviewStatus.CONFIRMED if total else FieldReviewStatus.MISSING,
                ),
            ]
        ),
    )


def test_build_column_headers_unique(schema: ExtractionSchema) -> None:
    headers = build_column_headers(schema)
    assert len(headers) == len(set(headers))
    assert headers[0] == "document_id"
    assert "Vendor" in headers
    assert "Vendor (review)" in headers


def test_preview_rows_includes_validated_and_review_required(schema: ExtractionSchema) -> None:
    docs = [
        _doc(status=ProcessingStatus.READY_FOR_EXPORT),
        _doc(status=ProcessingStatus.REVIEW_REQUIRED, total=None, review=FieldReviewStatus.MISSING),
    ]
    rows = preview_rows(docs, schema)
    assert len(rows) == 2
    assert rows[0]["export_review_status"] == "validated"
    assert rows[1]["export_review_status"] == "needs_review"
    assert rows[0]["Vendor"] == "Café Müller"


def test_export_new_workbook(tmp_path: Path, settings: Settings, schema: ExtractionSchema) -> None:
    target = tmp_path / "out.xlsx"
    doc = _doc(status=ProcessingStatus.READY_FOR_EXPORT)
    result = export_documents(target, [doc], schema, settings, mode="new")
    assert result.rows_written == 1
    assert target.exists()
    assert doc.status == ProcessingStatus.EXPORTED

    wb = load_workbook(target)
    sheet = wb[settings.excel_default_sheet]
    assert sheet.cell(1, 1).value == "document_id"
    assert sheet.cell(2, 2).value == "scan.png"
    assert "api" not in str(sheet.cell(2, 1).value).lower()


def test_refuse_overwrite_without_confirmation(
    tmp_path: Path, settings: Settings, schema: ExtractionSchema
) -> None:
    target = tmp_path / "out.xlsx"
    export_documents(target, [_doc(status=ProcessingStatus.READY_FOR_EXPORT)], schema, settings)
    with pytest.raises(ExcelExportError, match="overwrite"):
        export_documents(
            target,
            [_doc(status=ProcessingStatus.READY_FOR_EXPORT)],
            schema,
            settings,
            mode="new",
        )


def test_append_preserves_existing_rows(tmp_path: Path, settings: Settings, schema: ExtractionSchema) -> None:
    target = tmp_path / "out.xlsx"
    first = _doc(status=ProcessingStatus.READY_FOR_EXPORT)
    export_documents(target, [first], schema, settings, mode="new", mark_exported=False)

    wb = load_workbook(target)
    sheet = wb[settings.excel_default_sheet]
    sheet.cell(row=2, column=3, value="manual_keep")
    wb.save(target)

    second = _doc(status=ProcessingStatus.REVIEW_REQUIRED)
    export_documents(target, [second], schema, settings, mode="append", mark_exported=True)
    wb = load_workbook(target)
    sheet = wb[settings.excel_default_sheet]
    assert sheet.max_row == 3
    assert sheet.cell(2, 3).value == "manual_keep"
    assert sheet.cell(3, 2).value == "scan.png"


def test_append_rejects_invalid_workbook(tmp_path: Path, settings: Settings, schema: ExtractionSchema) -> None:
    target = tmp_path / "bad.xlsx"
    target.write_text("not an excel file", encoding="utf-8")
    with pytest.raises(ExcelExportError, match="Cannot append"):
        export_documents(
            target,
            [_doc(status=ProcessingStatus.READY_FOR_EXPORT)],
            schema,
            settings,
            mode="append",
        )


def test_append_rejects_mismatched_headers(tmp_path: Path, settings: Settings) -> None:
    target = tmp_path / "out.xlsx"
    schema_a = ExtractionSchema(
        name="a",
        fields=[ExtractionFieldDefinition(key="one", label="One")],
    )
    schema_b = ExtractionSchema(
        name="b",
        fields=[ExtractionFieldDefinition(key="two", label="Two")],
    )
    export_documents(target, [_doc(status=ProcessingStatus.READY_FOR_EXPORT)], schema_a, settings)
    with pytest.raises(ExcelExportError, match="headers"):
        export_documents(
            target,
            [_doc(status=ProcessingStatus.READY_FOR_EXPORT)],
            schema_b,
            settings,
            mode="append",
        )
